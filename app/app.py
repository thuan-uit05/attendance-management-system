import os
from flask import Flask, render_template, request, redirect, url_for, session, flash
from datetime import datetime
import mysql.connector
from dotenv import load_dotenv
from functools import wraps
from datetime import datetime
from flask import make_response
import openpyxl
from io import BytesIO

# Load biến môi trường
load_dotenv()

# Khởi tạo Flask với đường dẫn tuyệt đối tới templates
app = Flask(__name__,
            template_folder=os.path.abspath('templates'),
            static_folder=os.path.abspath('static'))

# Cấu hình bảo mật
app.secret_key = os.getenv('SECRET_KEY', 'your-secret-key-here')

# Debug: Kiểm tra đường dẫn
print(f"🛠 Đường dẫn làm việc hiện tại: {os.getcwd()}")
print(f"🛠 Đường dẫn templates: {app.template_folder}")
print(f"🛠 File trong templates: {os.listdir(app.template_folder)}")

# Cấu hình database
DB_CONFIG = {
    'host': os.getenv('DB_HOST', 'localhost'),
    'user': os.getenv('DB_USER', 'thith'),
    'password': os.getenv('DB_PASSWORD', '15112004'),
    'database': os.getenv('DB_NAME', 'Diemdanh'),
    'charset': 'utf8mb4'
}

def get_db_connection():
    """Tạo kết nối database với tự động đóng khi gặp lỗi"""
    try:
        conn = mysql.connector.connect(**DB_CONFIG)
        print("✅ Kết nối database thành công")
        return conn
    except mysql.connector.Error as err:
        print(f"❌ Lỗi kết nối database: {err}")
        flash('Lỗi kết nối cơ sở dữ liệu', 'danger')
        return None

# Decorator kiểm tra đăng nhập
def lecturer_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'lecturer_logged_in' not in session:
            flash('Vui lòng đăng nhập để tiếp tục', 'warning')
            return redirect(url_for('lecturer_login'))
        return f(*args, **kwargs)
    return decorated_function

# ============= ROUTES =============
@app.route('/')
def home():
    """Trang chủ hệ thống"""
    try:
        return render_template('home.html')
    except Exception as e:
        print(f"❌ Lỗi render template: {e}")
        return "Lỗi hệ thống: Không tải được trang chủ", 500

@app.route('/lecturer/login', methods=['GET', 'POST'])
def lecturer_login():
    """Xử lý đăng nhập giảng viên"""
    if request.method == 'POST':
        ma_giang_vien = request.form.get('ma_giang_vien')
        mat_khau = request.form.get('mat_khau')

        conn = get_db_connection()
        if conn:
            cursor = conn.cursor(dictionary=True)
            try:
                cursor.execute(
                    "SELECT * FROM GIANG_VIEN WHERE Ma_giang_vien = %s AND Mat_khau = %s",
                    (ma_giang_vien, mat_khau)
                )
                lecturer = cursor.fetchone()

                if lecturer:
                    session.update({
                        'lecturer_logged_in': True,
                        'ma_giang_vien': lecturer['Ma_giang_vien'],
                        'ho_ten': lecturer['Ho_ten']
                    })
                    flash('Đăng nhập thành công!', 'success')
                    return redirect(url_for('lecturer_dashboard'))
                else:
                    flash('Sai mã giảng viên hoặc mật khẩu', 'danger')
            except mysql.connector.Error as err:
                flash('Lỗi hệ thống, vui lòng thử lại sau', 'danger')
                print(f"Database error: {err}")
            finally:
                cursor.close()
                conn.close()

    return render_template('lecturer_login.html')

@app.route('/dashboard', endpoint='lecturer_dashboard')
@lecturer_required
def dashboard():
    """Trang dashboard giảng viên"""
    today = datetime.now().strftime('%Y-%m-%d')
    conn = get_db_connection()
    classes = []

    if conn:
        cursor = conn.cursor(dictionary=True)
        try:
            cursor.execute("""
                SELECT * FROM LOP_HOC 
                WHERE Ma_giang_vien = %s AND DATE(Thoi_gian_bat_dau) = %s
                ORDER BY Thoi_gian_bat_dau
            """, (session['ma_giang_vien'], today))
            classes = cursor.fetchall()
        except mysql.connector.Error as err:
            flash('Lỗi khi tải danh sách lớp', 'danger')
            print(f"Database error: {err}")
        finally:
            cursor.close()
            conn.close()

    return render_template('lecturer_dashboard.html',
                         classes=classes,
                         today=datetime.now().strftime('%d/%m/%Y'))

@app.route('/manual-attendance', methods=['GET', 'POST'], endpoint='manual_attendance')
@lecturer_required
def manual_attendance():
    """Trang điểm danh thủ công"""
    conn = get_db_connection()
    classes = []

    if conn:
        cursor = conn.cursor(dictionary=True)
        try:
            # Lấy tất cả lớp học của giảng viên này
            cursor.execute("""
                SELECT * FROM LOP_HOC 
                WHERE Ma_giang_vien = %s
                ORDER BY Thoi_gian_bat_dau DESC
            """, (session['ma_giang_vien'],))
            classes = cursor.fetchall()

            if request.method == 'POST':
                ma_lop = request.form.get('ma_lop')
                mssv = request.form.get('mssv')

                # Kiểm tra sinh viên có tồn tại không
                cursor.execute("SELECT * FROM SINH_VIEN WHERE MSSV = %s", (mssv,))
                student = cursor.fetchone()

                if not student:
                    flash('Không tìm thấy sinh viên với MSSV này', 'danger')
                    return render_template('manual_attendance.html', classes=classes)

                # Ghi nhận điểm danh
                cursor.execute("""
                    INSERT INTO DIEM_DANH (Ma_lop, MSSV, Thoi_gian_diem_danh, Trang_thai)
                    VALUES (%s, %s, NOW(), 'Có mặt')
                """, (ma_lop, mssv))
                conn.commit()
                flash('Điểm danh thành công!', 'success')

        except mysql.connector.Error as err:
            flash('Lỗi hệ thống khi xử lý điểm danh', 'danger')
            print(f"Database error: {err}")
        finally:
            cursor.close()
            conn.close()

    return render_template('manual_attendance.html', classes=classes)


# Thêm vào phần route trong app.py (sau route manual_attendance)

@app.route('/attendance-history')
@lecturer_required
def attendance_history():
    """Trang lịch sử điểm danh"""
    filters = {
        'ma_lop': request.args.get('ma_lop', 'all'),
        'start_date': request.args.get('start_date', ''),
        'end_date': request.args.get('end_date', '')
    }

    conn = get_db_connection()
    attendance_history = []
    classes = []

    if conn:
        cursor = conn.cursor(dictionary=True)
        try:
            # Lấy danh sách lớp học của giảng viên
            cursor.execute("""
                SELECT Ma_lop, Ten_mon_hoc 
                FROM LOP_HOC 
                WHERE Ma_giang_vien = %s
                ORDER BY Ten_mon_hoc
            """, (session['ma_giang_vien'],))
            classes = cursor.fetchall()

            # Xây dựng query lọc
            query = """
                SELECT dd.*, lh.Ten_mon_hoc, sv.Ho_ten, sv.MSSV
                FROM DIEM_DANH dd
                JOIN LOP_HOC lh ON dd.Ma_lop = lh.Ma_lop
                JOIN SINH_VIEN sv ON dd.MSSV = sv.MSSV
                WHERE lh.Ma_giang_vien = %s
            """
            params = [session['ma_giang_vien']]

            if filters['ma_lop'] != 'all':
                query += " AND dd.Ma_lop = %s"
                params.append(filters['ma_lop'])

            if filters['start_date']:
                query += " AND DATE(dd.Thoi_gian_diem_danh) >= %s"
                params.append(filters['start_date'])

            if filters['end_date']:
                query += " AND DATE(dd.Thoi_gian_diem_danh) <= %s"
                params.append(filters['end_date'])

            query += " ORDER BY dd.Thoi_gian_diem_danh DESC LIMIT 100"

            cursor.execute(query, params)
            attendance_history = cursor.fetchall()

        except mysql.connector.Error as err:
            flash('Lỗi khi tải lịch sử điểm danh', 'danger')
            print(f"Database error: {err}")
        finally:
            cursor.close()
            conn.close()

    return render_template('attendance_history.html',
                           attendance_history=attendance_history,
                           classes=classes,
                           filters=filters)


@app.route('/export-attendance')
@lecturer_required
def export_attendance():
    """Xuất Excel lịch sử điểm danh"""
    filters = {
        'ma_lop': request.args.get('ma_lop', 'all'),
        'start_date': request.args.get('start_date', ''),
        'end_date': request.args.get('end_date', '')
    }

    conn = get_db_connection()
    attendance_history = []

    if conn:
        cursor = conn.cursor(dictionary=True)
        try:
            # Xây dựng query lọc
            query = """
                SELECT dd.*, lh.Ten_mon_hoc, sv.Ho_ten, sv.MSSV
                FROM DIEM_DANH dd
                JOIN LOP_HOC lh ON dd.Ma_lop = lh.Ma_lop
                JOIN SINH_VIEN sv ON dd.MSSV = sv.MSSV
                WHERE lh.Ma_giang_vien = %s
            """
            params = [session['ma_giang_vien']]

            if filters['ma_lop'] != 'all':
                query += " AND dd.Ma_lop = %s"
                params.append(filters['ma_lop'])

            if filters['start_date']:
                query += " AND DATE(dd.Thoi_gian_diem_danh) >= %s"
                params.append(filters['start_date'])

            if filters['end_date']:
                query += " AND DATE(dd.Thoi_gian_diem_danh) <= %s"
                params.append(filters['end_date'])

            query += " ORDER BY dd.Thoi_gian_diem_danh DESC"

            cursor.execute(query, params)
            attendance_history = cursor.fetchall()

        except mysql.connector.Error as err:
            flash('Lỗi khi tải lịch sử điểm danh', 'danger')
            print(f"Database error: {err}")
        finally:
            cursor.close()
            conn.close()

    # Tạo file Excel
    output = BytesIO()
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Lịch sử điểm danh"

    # Thêm header
    headers = ["STT", "Thời gian", "Lớp học", "Sinh viên", "Trạng thái"]
    sheet.append(headers)

    # Thêm dữ liệu
    for idx, record in enumerate(attendance_history, 1):
        row = [
            idx,
            record['Thoi_gian_diem_danh'].strftime('%d/%m/%Y %H:%M'),
            f"{record['Ten_mon_hoc']} ({record['Ma_lop']})",
            f"{record['Ho_ten']} ({record['MSSV']})",
            record['Trang_thai']
        ]
        sheet.append(row)

    workbook.save(output)
    output.seek(0)

    response = make_response(output.getvalue())
    response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    response.headers['Content-Disposition'] = 'attachment; filename=lich_su_diem_danh.xlsx'
    return response

# Trong h�m class_detail(ma_lop) c?a app.py
@app.route('/class/<ma_lop>')
@lecturer_required
def class_detail(ma_lop):
    """Chi ti?t l?p h?c v� �i?m danh"""
    conn = get_db_connection()
    class_info = None
    registered_students = []
    attendance_list = []

    if conn:
        cursor = conn.cursor(dictionary=True)
        try:
            # L?y th�ng tin l?p h?c
            cursor.execute("""
                SELECT * FROM LOP_HOC 
                WHERE Ma_lop = %s AND Ma_giang_vien = %s
            """, (ma_lop, session['ma_giang_vien']))
            class_info = cursor.fetchone()

            if not class_info:
                flash('Kh�ng t?m th?y l?p h?c', 'danger')
                return redirect(url_for('lecturer_dashboard'))

            # L?y danh s�ch sinh vi�n �? ��ng k? v� tr?ng th�i �i?m danh
            cursor.execute("""
                SELECT sv.MSSV, sv.Ho_ten, 
                       CASE WHEN dd.MSSV IS NOT NULL THEN 1 ELSE 0 END as attendance_status
                FROM SINH_VIEN sv
                JOIN DANG_KY dk ON sv.MSSV = dk.MSSV
                LEFT JOIN DIEM_DANH dd ON sv.MSSV = dd.MSSV AND dd.Ma_lop = %s
                WHERE dk.Ma_lop = %s
                GROUP BY sv.MSSV, sv.Ho_ten
            """, (ma_lop, ma_lop))
            registered_students = cursor.fetchall()

            # L?y l?ch s? �i?m danh (10 b?n ghi g?n nh?t)
            cursor.execute("""
                SELECT dd.*, sv.Ho_ten 
                FROM DIEM_DANH dd
                JOIN SINH_VIEN sv ON dd.MSSV = sv.MSSV
                WHERE dd.Ma_lop = %s
                ORDER BY dd.Thoi_gian_diem_danh DESC
                LIMIT 10
            """, (ma_lop,))
            attendance_list = cursor.fetchall()

        except mysql.connector.Error as err:
            flash('L?i khi t?i th�ng tin l?p h?c', 'danger')
            print(f"Database error: {err}")
        finally:
            cursor.close()
            conn.close()

    return render_template('class_detail.html',
                         class_info=class_info,
                         registered_students=registered_students,
                         attendance_list=attendance_list)
@app.route('/warnings')
@lecturer_required
def view_warnings():
    """Trang quản lý cảnh báo"""
    conn = get_db_connection()
    warnings = []

    if conn:
        cursor = conn.cursor(dictionary=True)
        try:
            cursor.execute("""
                SELECT cb.*, dd.MSSV, sv.Ho_ten, lh.Ten_mon_hoc
                FROM CANH_BAO_GIAN_LAN cb
                JOIN DIEM_DANH dd ON cb.Ma_diem_danh = dd.Ma_diem_danh
                JOIN SINH_VIEN sv ON dd.MSSV = sv.MSSV
                JOIN LOP_HOC lh ON dd.Ma_lop = lh.Ma_lop
                WHERE lh.Ma_giang_vien = %s
                ORDER BY cb.Thoi_gian_phat_hien DESC
            """, (session['ma_giang_vien'],))
            warnings = cursor.fetchall()
        except mysql.connector.Error as err:
            flash('Lỗi khi tải cảnh báo', 'danger')
            print(f"Database error: {err}")
        finally:
            cursor.close()
            conn.close()

    return render_template('warnings.html', warnings=warnings)

@app.route('/logout')
def lecturer_logout():
    """Xử lý đăng xuất"""
    session.clear()
    flash('Bạn đã đăng xuất thành công', 'success')
    return redirect(url_for('home'))

@app.route('/data-management')
@lecturer_required
def data_management():
    """Trang quản lý dữ liệu"""
    tables = [
        {'name': 'SINH_VIEN', 'display': 'Sinh viên'},
        {'name': 'GIANG_VIEN', 'display': 'Giảng viên'},
        {'name': 'LOP_HOC', 'display': 'Lớp học'},
        {'name': 'THE_RFID', 'display': 'Thẻ RFID'},
        {'name': 'DANG_KY', 'display': 'Đăng ký lớp'},
    ]
    return render_template('data_management.html', tables=tables)

@app.route('/data-management/<table_name>')
@lecturer_required
def view_table(table_name):
    """Xem dữ liệu trong bảng"""
    if table_name not in ['SINH_VIEN', 'GIANG_VIEN', 'LOP_HOC', 'THE_RFID', 'DANG_KY']:
        flash('Bảng không tồn tại hoặc không được phép truy cập', 'danger')
        return redirect(url_for('data_management'))

    conn = get_db_connection()
    if conn:
        cursor = conn.cursor(dictionary=True)
        try:
            # Lấy dữ liệu bảng
            cursor.execute(f"SELECT * FROM {table_name} LIMIT 100")
            data = cursor.fetchall()

            # Lấy thông tin cột
            cursor.execute(f"SHOW COLUMNS FROM {table_name}")
            columns_info = cursor.fetchall()
            columns = [col['Field'] for col in columns_info]

            return render_template('view_table.html',
                                 table_name=table_name,
                                 display_name=table_name.replace('_', ' '),
                                 columns=columns,
                                 data=data)
        except mysql.connector.Error as err:
            flash(f'Lỗi khi truy vấn bảng: {err}', 'danger')
        finally:
            cursor.close()
            conn.close()

    return redirect(url_for('data_management'))


@app.route('/data-management/<table_name>/add', methods=['GET', 'POST'])
@lecturer_required
def add_record(table_name):
    """Thêm bản ghi mới"""
    if table_name not in ['SINH_VIEN', 'GIANG_VIEN', 'LOP_HOC', 'THE_RFID', 'DANG_KY']:
        flash('Bảng không tồn tại hoặc không được phép truy cập', 'danger')
        return redirect(url_for('data_management'))

    conn = get_db_connection()
    if conn:
        cursor = conn.cursor(dictionary=True)
        try:
            # Lấy thông tin cột
            cursor.execute(f"SHOW COLUMNS FROM {table_name}")
            columns_info = cursor.fetchall()

            if request.method == 'POST':
                # Xử lý thêm dữ liệu
                column_names = []
                values = []
                for col in columns_info:
                    if col['Field'] == 'id' or col['Extra'] == 'auto_increment':
                        continue

                    value = request.form.get(col['Field'], '')

                    # Xử lý đặc biệt cho trường ngày tháng
                    if col['Type'].startswith('date'):
                        if value:
                            try:
                                # Chuyển đổi từ định dạng HTML date (YYYY-MM-DD) sang MySQL date
                                value = datetime.strptime(value, '%Y-%m-%d').date()
                            except ValueError:
                                flash(f'Định dạng ngày không hợp lệ cho trường {col["Field"]}', 'danger')
                                return render_template('add_record.html',
                                                       table_name=table_name,
                                                       display_name=table_name.replace('_', ' '),
                                                       columns=columns_info)
                        elif col['Null'] == 'YES':
                            value = None
                        else:
                            flash(f'Trường {col["Field"]} là bắt buộc', 'danger')
                            return render_template('add_record.html',
                                                   table_name=table_name,
                                                   display_name=table_name.replace('_', ' '),
                                                   columns=columns_info)

                    column_names.append(col['Field'])
                    values.append(value)

                placeholders = ', '.join(['%s'] * len(values))
                query = f"INSERT INTO {table_name} ({', '.join(column_names)}) VALUES ({placeholders})"

                cursor.execute(query, values)
                conn.commit()
                flash('Thêm bản ghi thành công!', 'success')
                return redirect(url_for('view_table', table_name=table_name))

            return render_template('add_record.html',
                                   table_name=table_name,
                                   display_name=table_name.replace('_', ' '),
                                   columns=columns_info)
        except mysql.connector.Error as err:
            flash(f'Lỗi khi thêm bản ghi: {err}', 'danger')
            conn.rollback()
        finally:
            cursor.close()
            conn.close()

    return redirect(url_for('view_table', table_name=table_name))


@app.route('/data-management/<table_name>/edit/<id>', methods=['GET', 'POST'])
@lecturer_required
def edit_record(table_name, id):
    """S?a b?n ghi"""
    if table_name not in ['SINH_VIEN', 'GIANG_VIEN', 'LOP_HOC', 'THE_RFID', 'DANG_KY']:
        flash('B?ng không t?n t?i ho?c không đư?c phép truy c?p', 'danger')
        return redirect(url_for('data_management'))

    conn = get_db_connection()
    if conn:
        cursor = conn.cursor(dictionary=True)
        try:
            # L?y thông tin c?t
            cursor.execute(f"SHOW COLUMNS FROM {table_name}")
            columns_info = cursor.fetchall()

            # L?y b?n ghi hi?n t?i
            primary_key = get_primary_key(table_name)
            cursor.execute(f"SELECT * FROM {table_name} WHERE {primary_key} = %s", (id,))
            record = cursor.fetchone()

            if not record:
                flash('B?n ghi không t?n t?i', 'danger')
                return redirect(url_for('view_table', table_name=table_name))

            if request.method == 'POST':
                # X? l? c?p nh?t d? li?u
                updates = []
                values = []
                for col in columns_info:
                    if col['Field'] == primary_key:
                        continue

                    field = col['Field']
                    value = request.form.get(field, '')

                    # X? l? đ?c bi?t cho trư?ng ngày tháng
                    if col['Type'].startswith('date'):
                        if value:
                            try:
                                # Chuy?n đ?i t? đ?nh d?ng HTML date (YYYY-MM-DD) sang MySQL date
                                value = datetime.strptime(value, '%Y-%m-%d').date()
                            except ValueError:
                                flash(f'Đ?nh d?ng ngày không h?p l? cho trư?ng {field}', 'danger')
                                return render_template('edit_record.html',
                                                       table_name=table_name,
                                                       display_name=table_name.replace('_', ' '),
                                                       columns=columns_info,
                                                       record=record)
                        elif col['Null'] == 'YES':
                            value = None
                        else:
                            flash(f'Trư?ng {field} là b?t bu?c', 'danger')
                            return render_template('edit_record.html',
                                                   table_name=table_name,
                                                   display_name=table_name.replace('_', ' '),
                                                   columns=columns_info,
                                                   record=record)

                    updates.append(f"{field} = %s")
                    values.append(value)

                values.append(id)  # Thêm ID vào cu?i cho WHERE
                query = f"UPDATE {table_name} SET {', '.join(updates)} WHERE {primary_key} = %s"

                cursor.execute(query, values)
                conn.commit()
                flash('C?p nh?t b?n ghi thành công!', 'success')
                return redirect(url_for('view_table', table_name=table_name))

            return render_template('edit_record.html',
                                   table_name=table_name,
                                   display_name=table_name.replace('_', ' '),
                                   columns=columns_info,
                                   record=record)
        except mysql.connector.Error as err:
            flash(f'L?i khi c?p nh?t b?n ghi: {err}', 'danger')
            conn.rollback()
        finally:
            cursor.close()
            conn.close()

    return redirect(url_for('view_table', table_name=table_name))

@app.route('/data-management/<table_name>/delete/<id>')
@lecturer_required
def delete_record(table_name, id):
    """Xóa bản ghi"""
    if table_name not in ['SINH_VIEN', 'GIANG_VIEN', 'LOP_HOC', 'THE_RFID', 'DANG_KY']:
        flash('Bảng không tồn tại hoặc không được phép truy cập', 'danger')
        return redirect(url_for('data_management'))

    conn = get_db_connection()
    if conn:
        cursor = conn.cursor()
        try:
            primary_key = get_primary_key(table_name)
            cursor.execute(f"DELETE FROM {table_name} WHERE {primary_key} = %s", (id,))
            conn.commit()
            flash('Xóa bản ghi thành công!', 'success')
        except mysql.connector.Error as err:
            flash(f'Lỗi khi xóa bản ghi: {err}', 'danger')
            conn.rollback()
        finally:
            cursor.close()
            conn.close()

    return redirect(url_for('view_table', table_name=table_name))

def get_primary_key(table_name):
    """Hàm helper để lấy tên khóa chính của bảng"""
    primary_keys = {
        'SINH_VIEN': 'MSSV',
        'GIANG_VIEN': 'Ma_giang_vien',
        'LOP_HOC': 'Ma_lop',
        'THE_RFID': 'Ma_the',
        'DANG_KY': 'Ma_dang_ky'
    }
    return primary_keys.get(table_name, 'id')

if __name__ == '__main__':
    # Kiểm tra kết nối database khi khởi động
    test_conn = get_db_connection()
    if test_conn:
        test_conn.close()

    app.run(host='0.0.0.0', port=5000, debug=True)